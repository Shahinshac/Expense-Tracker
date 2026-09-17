import io
import csv
from datetime import date
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models.models import Expense

def generate_csv_export(expenses: List[Expense]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Date", "Time", "Type", "Category", "Amount (INR)", "Payment Method", "Description", "Note"])
    for exp in expenses:
        cat_name = exp.category.name if exp.category else "Uncategorized"
        amt = f"{exp.amount_paise / 100:.2f}"
        writer.writerow([
            exp.date.isoformat(),
            exp.time or "",
            "Expense",
            cat_name,
            amt,
            exp.payment_method or "",
            exp.description or "",
            exp.note or ""
        ])
    return output.getvalue()

def generate_excel_export(expenses: List[Expense], start_date: date, end_date: date) -> io.BytesIO:
    wb = Workbook()
    
    # Sheet 1: Transactions
    ws1 = wb.active
    ws1.title = "Transactions"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = Font(bold=True, color="FFFFFF")
    primary_header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    
    headers = ["Date", "Time", "Category", "Group", "Amount (₹)", "Payment Method", "Description", "Note"]
    ws1.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = primary_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_paise = 0
    category_totals = {}

    for exp in expenses:
        cat_name = exp.category.name if exp.category else "Other"
        cat_group = exp.category.group if exp.category else "Other"
        amt_rupees = exp.amount_paise / 100.0
        total_paise += exp.amount_paise
        category_totals[cat_name] = category_totals.get(cat_name, 0) + exp.amount_paise

        ws1.append([
            exp.date.isoformat(),
            exp.time or "",
            cat_name,
            cat_group,
            amt_rupees,
            exp.payment_method or "",
            exp.description or "",
            exp.note or ""
        ])

    # Format numeric column
    for row in range(2, len(expenses) + 2):
        ws1.cell(row=row, column=5).number_format = '₹#,##0.00'

    # Auto adjust column widths
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Sheet 2: Summary & Category Totals
    ws2 = wb.create_sheet(title="Summary")
    ws2.append(["Period", f"{start_date.isoformat()} to {end_date.isoformat()}"])
    ws2.append(["Total Transactions", len(expenses)])
    ws2.append(["Total Spending (₹)", total_paise / 100.0])
    ws2.cell(row=3, column=2).number_format = '₹#,##0.00'
    ws2.append([])
    
    ws2.append(["Category", "Total Spent (₹)", "Percentage"])
    for col_num in range(1, 4):
        c = ws2.cell(row=5, column=col_num)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")

    row_idx = 6
    for cat_name, amt_p in sorted(category_totals.items(), key=lambda x: x[1], reverse=True):
        pct = (amt_p / total_paise * 100) if total_paise > 0 else 0
        ws2.append([cat_name, amt_p / 100.0, f"{pct:.1f}%"])
        ws2.cell(row=row_idx, column=2).number_format = '₹#,##0.00'
        row_idx += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 16)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
